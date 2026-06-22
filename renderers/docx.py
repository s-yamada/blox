"""
blox - Word (.docx) レンダラー
同じYAMLスキーマからWord文書を生成する。
レイアウト（grid/col）はコンテンツを縦積みで表現。
"""
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _add_page_break(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_break(docx_break_type())


def docx_break_type():
    from docx.enum.text import WD_BREAK
    return WD_BREAK.PAGE


def _set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def _para_text(doc, text, bold=False, size=None, color=None, align=None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return p


# ============================================================
# セル描画
# ============================================================

def _img_placeholder_docx(doc, label, height_mm=35):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = 'Table Grid'
    cell = tbl.rows[0].cells[0]
    # 行の高さを固定
    tr = tbl.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    trH = OxmlElement('w:trHeight')
    trH.set(qn('w:val'), str(int(height_mm / 25.4 * 72 * 20)))  # mm → twips
    trH.set(qn('w:hRule'), 'exact')
    trPr.append(trH)
    _set_cell_bg(cell, 'CCCCCC')
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(f'📷  {label}')
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def _cell_img_docx(doc, cell):
    label = cell.get('image_label', '写真') if isinstance(cell, dict) else str(cell)
    _img_placeholder_docx(doc, label)


def _cell_text_docx(doc, cell):
    if cell.get('title'):
        _para_text(doc, cell['title'], bold=True, size=11)
    if cell.get('text'):
        for line in cell['text'].split('\n'):
            doc.add_paragraph(line)


def _cell_list_docx(doc, cell):
    if cell.get('title'):
        _para_text(doc, cell['title'], bold=True, size=11)
    items = cell.get('items', [])
    list_type = cell.get('list_type', 'ul')
    for i, item in enumerate(items):
        if list_type == 'ol':
            doc.add_paragraph(f'{i+1}. {item}')
        elif list_type == 'dl':
            sep = '：' if '：' in item else ':'
            if sep in item:
                term, desc = item.split(sep, 1)
                p = doc.add_paragraph()
                run = p.add_run(term.strip() + ': ')
                run.bold = True
                p.add_run(desc.strip())
            else:
                _para_text(doc, item, bold=True)
        else:
            doc.add_paragraph(f'• {item}')


def _cell_card_docx(doc, cell):
    label = cell.get('image_label', '')
    if label:
        _img_placeholder_docx(doc, label, height_mm=28)
    if cell.get('title'):
        _para_text(doc, cell['title'], bold=True)
    if cell.get('text'):
        doc.add_paragraph(cell['text'])


def _cell_table_docx(doc, cell):
    if cell.get('title'):
        _para_text(doc, cell['title'], bold=True, size=11)
    headers = cell.get('headers', [])
    items   = cell.get('items', [])
    if not items and not headers:
        return
    n_cols = max(
        len(headers) if headers else 0,
        max((len(r) if isinstance(r, list) else 1) for r in items) if items else 1,
    )
    tbl = doc.add_table(rows=0, cols=n_cols)
    tbl.style = 'Table Grid'
    if headers:
        row = tbl.add_row()
        for i, h in enumerate(headers[:n_cols]):
            row.cells[i].text = str(h)
            _set_cell_bg(row.cells[i], 'CCCCCC')
    for data_row in items:
        row = tbl.add_row()
        vals = data_row if isinstance(data_row, list) else [str(data_row)]
        for i, v in enumerate(vals[:n_cols]):
            row.cells[i].text = str(v)


def _cell_form_docx(doc, cell):
    items = cell.get('items', [])
    tbl = doc.add_table(rows=0, cols=2)
    tbl.style = 'Table Grid'
    for item in items:
        ftype = item.get('type', 'text')
        label = item.get('label', '')
        if ftype in ('submit', 'button'):
            row = tbl.add_row()
            row.cells[0].merge(row.cells[1])
            row.cells[0].text = f'[ {label} ]'
        elif ftype == 'buttons':
            row = tbl.add_row()
            row.cells[0].merge(row.cells[1])
            row.cells[0].text = '  '.join(f'[ {b.get("label","")} ]' for b in item.get('items', []))
        elif ftype == 'label':
            row = tbl.add_row()
            row.cells[0].text = label
            row.cells[1].text = item.get('value', '')
        elif ftype == 'select':
            opts = ' / '.join(item.get('options', []))
            row = tbl.add_row()
            row.cells[0].text = label
            row.cells[1].text = f'▾ {opts}'
        elif ftype == 'checkbox':
            row = tbl.add_row()
            row.cells[0].merge(row.cells[1])
            row.cells[0].text = f'☐ {item.get("text") or label}'
        elif ftype == 'radio':
            for o in item.get('options', []):
                row = tbl.add_row()
                row.cells[0].merge(row.cells[1])
                row.cells[0].text = f'◯ {o}'
        else:
            row = tbl.add_row()
            row.cells[0].text = label
            row.cells[1].text = item.get('placeholder', '')


def _dispatch_cell_docx(doc, cell, kind=None):
    k = kind or (cell.get('kind', 'text') if isinstance(cell, dict) else 'img')
    if k == 'img':    _cell_img_docx(doc, cell)
    elif k == 'text': _cell_text_docx(doc, cell)
    elif k == 'list': _cell_list_docx(doc, cell)
    elif k == 'table': _cell_table_docx(doc, cell)
    elif k == 'card': _cell_card_docx(doc, cell)
    elif k == 'form': _cell_form_docx(doc, cell)


# ============================================================
# ブロック描画
# ============================================================

def block_hero_docx(doc, block):
    label   = block.get('image_label', '写真')
    caption = block.get('caption', '')
    _img_placeholder_docx(doc, label, height_mm=50)
    if caption:
        _para_text(doc, caption, bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)


def block_stats_docx(doc, block):
    items = block.get('items', [])
    if not items:
        return
    tbl = doc.add_table(rows=2, cols=len(items))
    tbl.style = 'Table Grid'
    for i, s in enumerate(items):
        tbl.rows[0].cells[i].text = s['num']
        tbl.rows[1].cells[i].text = s['label']
        tbl.rows[0].cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        tbl.rows[1].cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER


def block_cta_docx(doc, block):
    text = block.get('text', 'お問い合わせはこちら')
    _para_text(doc, f'▶ {text}', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)


def block_steps_docx(doc, block):
    items = block.get('items', [])
    for i, s in enumerate(items):
        num   = s.get('num', f'{i+1:02d}')
        title = s.get('title', '')
        text  = s.get('text', '')
        p = doc.add_paragraph()
        run = p.add_run(f'{num}. {title}')
        run.bold = True
        if text:
            p.add_run(f' — {text}')


def block_grid_docx(doc, block):
    kind  = block.get('kind', 'card')
    items = block.get('items', [])
    for item in items:
        if isinstance(item, str):
            item = {'image_label': item}
        _dispatch_cell_docx(doc, {**item, 'kind': kind}, kind)
        doc.add_paragraph()


def block_col_docx(doc, block):
    rows_data = block.get('rows') or [{'cells': block.get('cells', [])}]
    for row in rows_data:
        for cell in row.get('cells', []):
            _dispatch_cell_docx(doc, cell)


def render_block_docx(doc, block):
    if block.get('title'):
        doc.add_heading(block['title'], level=2)
    btype = block.get('type', '')
    fn = globals().get(f'block_{btype}_docx')
    if fn:
        fn(doc, block)


def build_docx(conf, out_path):
    doc = Document()

    # A4縦
    section = doc.sections[0]
    section.page_width  = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = section.right_margin = Mm(20)
    section.top_margin  = section.bottom_margin = Mm(20)

    # 表紙
    if conf.get('title'):
        doc.add_heading(conf['title'], level=0)
        if conf.get('author'):
            _para_text(doc, conf['author'], align=WD_ALIGN_PARAGRAPH.CENTER)
        doc.add_page_break()

    pages = conf.get('pages', [])
    for i, page in enumerate(pages):
        doc.add_heading(f'No.{i+1:02d}  {page["title"]}', level=1)
        for block in page.get('blocks', []):
            render_block_docx(doc, block)
            doc.add_paragraph()
        if i < len(pages) - 1:
            doc.add_page_break()

    doc.save(out_path)
