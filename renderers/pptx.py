"""
blox - PPTXパターンライブラリ
各関数は (slide, y, ...) を受け取り、描画後の「次のy座標」を返す
"""
from pptx.util import Mm, Pt, Emu
from pptx.enum.text import PP_ALIGN
from lxml import etree
from renderers.pptx_styles import F, GAP, GAP_SM, MARGIN_L, CONTENT_W, FONT
import themes
from themes import C

_CELL_GAP = Mm(3)  # grid/col 内のセル間ギャップ

_EFFECT_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'
_EMPTY_EFFECT_LST = f'<a:effectLst xmlns:a="{_EFFECT_NS}"/>'


def _add_shape(slide, x, y, w, h, fill_color=None, line_color=None, line_w=None):
    shape = slide.shapes.add_shape(1, x, y, w, h)
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if line_color:
        shape.line.color.rgb = line_color
        if line_w:
            shape.line.width = line_w
    else:
        shape.line.fill.background()
    if not themes.shadow:
        spPr = shape._element.spPr
        for el in spPr.findall(f'{{{_EFFECT_NS}}}effectLst'):
            spPr.remove(el)
        spPr.append(etree.fromstring(_EMPTY_EFFECT_LST))
    return shape


def _add_textbox(slide, x, y, w, h, text, font_size, color, bold=False,
                 align=PP_ALIGN.LEFT, v_anchor=None, wrap=True):
    from pptx.enum.text import MSO_ANCHOR
    txb = slide.shapes.add_textbox(x, y, w, h)
    tf = txb.text_frame
    tf.word_wrap = wrap
    if v_anchor:
        tf.vertical_anchor = v_anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = str(text) if text is not None else ''
    run.font.size = font_size
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = FONT
    return txb


def _set_shape_text(shape, text, font_size, color, bold=False,
                    align=PP_ALIGN.LEFT, v_anchor=None):
    from pptx.enum.text import MSO_ANCHOR
    tf = shape.text_frame
    tf.word_wrap = True
    if v_anchor:
        tf.vertical_anchor = v_anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = str(text) if text is not None else ''
    run.font.size = font_size
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = FONT


# ============================================================
# セル描画ヘルパー（grid / col から呼び出す）
# ============================================================

def _draw_img_cell(slide, x, y, w, h, cell):
    label = cell.get('image_label', '写真') if isinstance(cell, dict) else str(cell)
    img = _add_shape(slide, x, y, w, h, fill_color=C['hero_bg'])
    _set_shape_text(img, label, F['small'], C['blue_dk'], align=PP_ALIGN.CENTER)


def _draw_text_cell(slide, x, y, w, h, cell):
    used_h = Emu(0)
    if cell.get('title'):
        ttl_h = Mm(9)
        _add_shape(slide, x, y + used_h, Mm(3), ttl_h, fill_color=C['accent'])
        ttl = _add_shape(slide, x + Mm(5), y + used_h, w - Mm(5), ttl_h,
                         fill_color=None)
        _set_shape_text(ttl, cell['title'], F['section'], C['navy'], bold=True)
        used_h += ttl_h + GAP_SM
    if cell.get('text'):
        body_h = h - used_h
        bod = _add_shape(slide, x, y + used_h, w, body_h, fill_color=None)
        _set_shape_text(bod, cell['text'], F['body'], C['blue_dk'])


def _draw_list_cell(slide, x, y, w, h, cell):
    items = cell.get('items', [])
    if not items:
        return
    list_type = cell.get('list_type', 'ul')
    cy, ch = y, h
    if cell.get('title'):
        ttl_h = Mm(8)
        _add_shape(slide, x, cy, Mm(3), ttl_h, fill_color=C['accent'])
        ttl = _add_shape(slide, x + Mm(5), cy, w - Mm(5), ttl_h, fill_color=None)
        _set_shape_text(ttl, cell['title'], F['section'], C['navy'], bold=True)
        cy += ttl_h + GAP_SM
        ch  = h - ttl_h - GAP_SM
    _add_shape(slide, x, cy, w, ch,
               fill_color=C['gray_lt'], line_color=C['divider'], line_w=Mm(0.2))
    item_h = ch / max(len(items), 1)
    for i, item in enumerate(items):
        prefix = f'{i+1}. ' if list_type == 'ol' else '▶ '
        iy = cy + item_h * i
        itm = _add_shape(slide, x + Mm(2), iy, w - Mm(4), item_h, fill_color=None)
        _set_shape_text(itm, f'{prefix}{item}', F['small'], C['blue_dk'])


def _draw_card_cell(slide, x, y, w, h, cell):
    img_ratio = 0.55
    img_h = Emu(int(h * img_ratio))
    _add_shape(slide, x, y, w, h, fill_color=C['white'],
               line_color=C['divider'], line_w=Mm(0.2))
    img = _add_shape(slide, x, y, w, img_h, fill_color=C['card_img'])
    _set_shape_text(img, cell.get('image_label', '写真'), F['small'],
                    C['blue_dk'], align=PP_ALIGN.CENTER)
    text_y = y + img_h
    if cell.get('title'):
        ttl = _add_shape(slide, x + Mm(2), text_y + Mm(2),
                         w - Mm(4), Mm(8), fill_color=None)
        _set_shape_text(ttl, cell['title'], F['small'], C['navy'], bold=True)
        text_y += Mm(10)
    if cell.get('text'):
        remaining = h - (text_y - y) - Mm(2)
        if remaining > Emu(0):
            txt = _add_shape(slide, x + Mm(2), text_y,
                             w - Mm(4), remaining, fill_color=None)
            _set_shape_text(txt, cell['text'], F['small'], C['blue_dk'])


def _draw_table_cell(slide, x, y, w, h, cell):
    headers = cell.get('headers', [])
    items   = cell.get('items', [])
    all_rows = (([headers] if headers else []) + list(items)) if (headers or items) else []
    if not all_rows:
        return
    cy, ch = y, h
    if cell.get('title'):
        ttl_h = Mm(8)
        _add_shape(slide, x, cy, Mm(3), ttl_h, fill_color=C['accent'])
        ttl = _add_shape(slide, x + Mm(5), cy, w - Mm(5), ttl_h, fill_color=None)
        _set_shape_text(ttl, cell['title'], F['section'], C['navy'], bold=True)
        cy += ttl_h + GAP_SM
        ch  = h - ttl_h - GAP_SM
    n_rows = len(all_rows)
    row_h  = ch / n_rows
    n_cols = max((len(r) if isinstance(r, list) else 1) for r in all_rows)
    col_w  = w / max(n_cols, 1)

    for ri, row in enumerate(all_rows):
        is_header = ri == 0 and bool(headers)
        ry = cy + row_h * ri
        bg = C['navy'] if is_header else (C['gray_lt'] if ri % 2 == 0 else C['white'])
        _add_shape(slide, x, ry, w, row_h, fill_color=bg,
                   line_color=C['divider'], line_w=Mm(0.15))
        vals = row if isinstance(row, list) else [str(row)]
        for ci, val in enumerate(vals):
            text_color = C['white'] if is_header else C['blue_dk']
            cx = x + col_w * ci
            tb = _add_shape(slide, cx + Mm(1), ry, col_w - Mm(2), row_h,
                            fill_color=None)
            _set_shape_text(tb, str(val), F['small'], text_color, bold=is_header)


def _draw_form_cell(slide, x, y, w, h, cell):
    items = cell.get('items', [])
    if not items:
        return
    label_w = Emu(int(w * 0.3))
    field_w = w - label_w - Mm(2)
    field_x = x + label_w + Mm(2)
    row_h = h / max(len(items), 1)

    for i, item in enumerate(items):
        iy = y + row_h * i
        ftype = item.get('type', 'text') if isinstance(item, dict) else 'text'
        label = item.get('label', '') if isinstance(item, dict) else str(item)

        if ftype in ('submit', 'button'):
            color = C['navy'] if ftype == 'submit' or item.get('color', 'primary') == 'primary' else C['gray_md']
            btn = _add_shape(slide, field_x, iy + Mm(1),
                             Emu(int(field_w * 0.5)), row_h - Mm(2),
                             fill_color=color)
            _set_shape_text(btn, label, F['small'], C['white'],
                            bold=True, align=PP_ALIGN.CENTER)
        elif ftype == 'buttons':
            sub = item.get('items', [])
            n = max(len(sub), 1)
            btn_w = Emu(int(field_w / n)) - Mm(2)
            for j, b in enumerate(sub):
                bc = C['navy'] if b.get('color', 'secondary') == 'primary' else C['gray_md']
                bx = field_x + (btn_w + Mm(2)) * j
                sb = _add_shape(slide, bx, iy + Mm(1), btn_w, row_h - Mm(2),
                                fill_color=bc)
                _set_shape_text(sb, b.get('label', ''), F['small'], C['white'],
                                bold=True, align=PP_ALIGN.CENTER)
        elif ftype == 'label':
            lbl = _add_shape(slide, x, iy, label_w, row_h, fill_color=None)
            _set_shape_text(lbl, label, F['small'], C['gray_md'])
            val = _add_shape(slide, field_x, iy, field_w, row_h, fill_color=None)
            _set_shape_text(val, item.get('value', ''), F['small'], C['blue_dk'])
        else:
            lbl = _add_shape(slide, x, iy, label_w, row_h, fill_color=None)
            _set_shape_text(lbl, label, F['small'], C['blue_dk'])

            if ftype in ('checkbox', 'radio'):
                prefix = '☑ ' if ftype == 'checkbox' else '◉ '
                options = item.get('options', [item.get('text', '')])
                opt_w = Emu(int(field_w / max(len(options), 1)))
                for j, opt in enumerate(options):
                    ob = _add_shape(slide, field_x + opt_w * j, iy,
                                    opt_w, row_h, fill_color=None)
                    _set_shape_text(ob, f'{prefix}{opt}', F['small'], C['blue_dk'])
            elif ftype == 'file':
                fld = _add_shape(slide, field_x, iy + Mm(1),
                                 field_w, row_h - Mm(2),
                                 fill_color=C['gray_lt'],
                                 line_color=C['divider'], line_w=Mm(0.2))
                _set_shape_text(fld, 'ファイルを選択…', F['small'], C['gray_md'])
            else:
                fld = _add_shape(slide, field_x, iy + Mm(1),
                                 field_w, row_h - Mm(2),
                                 fill_color=C['gray_lt'],
                                 line_color=C['divider'], line_w=Mm(0.2))
                placeholder = item.get('placeholder', '')
                if placeholder:
                    _set_shape_text(fld, placeholder, F['small'], C['gray_md'])


def _dispatch_cell(slide, x, y, w, h, cell):
    kind = cell.get('kind', 'text') if isinstance(cell, dict) else 'img'
    if kind == 'img':
        _draw_img_cell(slide, x, y, w, h, cell)
    elif kind == 'text':
        _draw_text_cell(slide, x, y, w, h, cell)
    elif kind == 'list':
        _draw_list_cell(slide, x, y, w, h, cell)
    elif kind == 'table':
        _draw_table_cell(slide, x, y, w, h, cell)
    elif kind == 'card':
        _draw_card_cell(slide, x, y, w, h, cell)
    elif kind == 'form':
        _draw_form_cell(slide, x, y, w, h, cell)


# ============================================================
# 表紙
# ============================================================
def add_cover(slide, title, author=''):
    """表紙スライド（ページ番号なし）"""
    from renderers.pptx_styles import SLIDE_W, SLIDE_H
    cover_h = Mm(130)
    _add_shape(slide, Mm(0), Mm(0), SLIDE_W, cover_h, fill_color=C['navy'])
    # タイトル
    ttl = _add_shape(slide, Mm(20), Mm(45), SLIDE_W - Mm(40), Mm(50),
                     fill_color=None)
    _set_shape_text(ttl, title, F['cover_title'], C['white'],
                    bold=True, align=PP_ALIGN.CENTER)
    # 著者
    if author:
        auth = _add_shape(slide, Mm(20), Mm(150), SLIDE_W - Mm(40), Mm(15),
                          fill_color=None)
        _set_shape_text(auth, author, F['cover_author'], C['blue_dk'],
                        align=PP_ALIGN.CENTER)


# ============================================================
# ヘッダー
# ============================================================
def add_header(slide, page_no, title):
    from renderers.pptx_styles import HEADER_H, BADGE_W, SLIDE_W
    x, y = Mm(0), Mm(0)
    w, h = SLIDE_W, HEADER_H

    _add_shape(slide, x, y, w, h, fill_color=C['navy'])

    badge = _add_shape(slide, MARGIN_L, Mm(2), BADGE_W, Mm(14),
                       fill_color=C['teal_lt'])
    _set_shape_text(badge, f'No.{page_no:02d}', F['badge'], C['navy'],
                    bold=True, align=PP_ALIGN.CENTER)

    tx = MARGIN_L + BADGE_W + Mm(3)
    title_box = _add_shape(slide, tx, Mm(2), w - tx - MARGIN_L, Mm(14),
                           fill_color=None)
    _set_shape_text(title_box, title, F['page_title'], C['white'], bold=True)

    return HEADER_H


# ============================================================
# スタンドアロンブロック
# ============================================================

def block_hero(slide, y, h=Mm(55), image_label='写真', caption=''):
    x, w = MARGIN_L, CONTENT_W
    img_h = h - (Mm(10) if caption else Emu(0))
    img = _add_shape(slide, x, y, w, img_h, fill_color=C['hero_bg'])
    _set_shape_text(img, image_label, F['body'], C['blue_dk'],
                    align=PP_ALIGN.CENTER)
    next_y = y + img_h
    if caption:
        cap = _add_shape(slide, x, next_y, w, Mm(10), fill_color=C['navy'])
        _set_shape_text(cap, caption, F['small'], C['white'],
                        align=PP_ALIGN.CENTER)
        next_y += Mm(10)
    return next_y + GAP


def block_stats(slide, y, stats, h=Mm(22)):
    x, w = MARGIN_L, CONTENT_W
    n = len(stats)
    cell_w = w // n
    _add_shape(slide, x, y, w, h, fill_color=C['banner_bg'])
    for i, s in enumerate(stats):
        cx = x + cell_w * i
        num_box = _add_shape(slide, cx, y, cell_w, Mm(13), fill_color=None)
        _set_shape_text(num_box, s['num'], F['stat_num'], C['navy'],
                        bold=True, align=PP_ALIGN.CENTER)
        lbl_box = _add_shape(slide, cx, y + Mm(13), cell_w, Mm(9),
                             fill_color=None)
        _set_shape_text(lbl_box, s['label'], F['stat_label'], C['blue_dk'],
                        align=PP_ALIGN.CENTER)
    return y + h + GAP


def block_cta(slide, y, text='お問い合わせはこちら', h=Mm(14)):
    x, w = MARGIN_L, CONTENT_W
    bar = _add_shape(slide, x, y, w, h, fill_color=C['navy'])
    _set_shape_text(bar, text, F['body'], C['white'],
                    bold=True, align=PP_ALIGN.CENTER)
    return y + h + GAP


def block_steps(slide, y, steps, h=Mm(30)):
    x, w = MARGIN_L, CONTENT_W
    n = len(steps)
    cell_w = (w - _CELL_GAP * (n - 1)) // n
    for i, step in enumerate(steps):
        cx = x + (cell_w + _CELL_GAP) * i
        num_h = Mm(8)
        num = _add_shape(slide, cx, y, cell_w, num_h, fill_color=C['navy'])
        _set_shape_text(num, str(step.get('num', f'{i+1:02d}')), F['small'],
                        C['white'], bold=True, align=PP_ALIGN.CENTER)
        ttl_h = Mm(8)
        ttl = _add_shape(slide, cx, y + num_h, cell_w, ttl_h,
                         fill_color=C['gray_lt'])
        _set_shape_text(ttl, step.get('title', ''), F['small'], C['navy'],
                        bold=True, align=PP_ALIGN.CENTER)
        body_h = h - num_h - ttl_h
        bod = _add_shape(slide, cx, y + num_h + ttl_h, cell_w, body_h,
                         fill_color=None)
        _set_shape_text(bod, step.get('text', ''), F['small'], C['blue_dk'])
    return y + h + GAP


# ============================================================
# レイアウトブロック
# ============================================================

def block_grid(slide, y, kind='card', cols=3, items=None, item_h=Mm(50)):
    """
    均質グリッド: 同じ kind のセルを cols 列で並べる。
    items が cols を超える場合は折り返して複数行になる。
    """
    if items is None:
        items = []
    x, w = MARGIN_L, CONTENT_W
    cell_w = (w - _CELL_GAP * (cols - 1)) // cols
    rows = max((len(items) + cols - 1) // cols, 1)
    total_h = item_h * rows + _CELL_GAP * (rows - 1)

    for i, item in enumerate(items):
        row = i // cols
        col_idx = i % cols
        cx = x + (cell_w + _CELL_GAP) * col_idx
        cy = y + (item_h + _CELL_GAP) * row
        if isinstance(item, str):
            item = {'image_label': item, 'kind': kind}
        _dispatch_cell(slide, cx, cy, cell_w, item_h,
                       {**item, 'kind': kind})

    return y + total_h + GAP


def block_col(slide, y, rows=None, total_cols=12, item_h=Mm(50)):
    """
    異種横並び（複数行対応）: rows に行リストを渡すと縦に積まれる。
    各行の cells に kind と span を指定する。
    """
    if not rows:
        return y + item_h + GAP
    available_w = CONTENT_W
    col_unit = available_w / total_cols
    curr_y = y
    for row_idx, row in enumerate(rows):
        cells = row.get('cells', [])
        n = len(cells)
        if not cells:
            curr_y += item_h + GAP_SM
            continue
        row_available_w = CONTENT_W - _CELL_GAP * (n - 1)
        row_col_unit = row_available_w / total_cols
        curr_x = MARGIN_L
        for i, cell in enumerate(cells):
            span = cell.get('span', 1)
            cell_w = Emu(int(row_col_unit * span))
            _dispatch_cell(slide, curr_x, curr_y, cell_w, item_h, cell)
            curr_x += cell_w + (_CELL_GAP if i < n - 1 else Emu(0))
        curr_y += item_h + GAP_SM
    return curr_y - GAP_SM + GAP


_SECTION_TITLE_H = Mm(7)


def render_block(slide, y, block):
    """ブロック種別を判定してPPTX描画関数を呼び出す。
    組み込みブロックは明示的にdispatch、未知のブロックはプラグイン関数を探す。
    プラグイン関数のシグネチャ: block_{type}(slide, y, block) -> next_y
    """
    if block.get('title'):
        _add_textbox(slide, MARGIN_L, y, CONTENT_W, _SECTION_TITLE_H,
                     block['title'], F['section'], C['navy'], bold=True)
        y += _SECTION_TITLE_H + GAP_SM
    btype = block.get('type', '')
    if btype == 'hero':
        return block_hero(slide, y,
            h=Mm(block.get('h_mm', 55)),
            image_label=block.get('image_label', '写真'),
            caption=block.get('caption', ''),
        )
    elif btype == 'stats':
        return block_stats(slide, y,
            stats=block.get('items', []),
            h=Mm(block.get('h_mm', 22)),
        )
    elif btype == 'cta':
        return block_cta(slide, y,
            text=block.get('text', 'お問い合わせはこちら'),
            h=Mm(block.get('h_mm', 14)),
        )
    elif btype == 'steps':
        return block_steps(slide, y,
            steps=block.get('items', []),
            h=Mm(block.get('h_mm', 30)),
        )
    elif btype == 'grid':
        return block_grid(slide, y,
            kind=block.get('kind', 'card'),
            cols=block.get('cols', 3),
            items=block.get('items', []),
            item_h=Mm(block.get('item_h_mm', 50)),
        )
    elif btype == 'col':
        rows = block.get('rows') or [{'cells': block.get('cells', [])}]
        return block_col(slide, y,
            rows=rows,
            total_cols=block.get('cols', 12),
            item_h=Mm(block.get('item_h_mm', block.get('h_mm', 50))),
        )
    fn = globals().get(f'block_{btype}')
    if fn:
        return fn(slide, y, block)
    import sys
    print(f'  ⚠ 未知のブロックタイプ: {btype}', file=sys.stderr)
    return y


def measure_block(block):
    """ブロックの高さ（EMU）を render_block() と同じロジックで概算する。"""
    btype = block.get('type', '')
    title_h = _SECTION_TITLE_H + GAP_SM if block.get('title') else Emu(0)
    _DEFAULTS = {'hero': 55, 'stats': 22, 'cta': 14, 'steps': 30}
    if btype == 'grid':
        cols    = block.get('cols', 3)
        n_items = len(block.get('items', []))
        n_rows  = max(1, -(-n_items // cols))  # ceil
        item_h  = Mm(block.get('item_h_mm', 50))
        content_h = item_h * n_rows + _CELL_GAP * (n_rows - 1)
    elif btype == 'col':
        rows_data = block.get('rows')
        if rows_data:
            n       = len(rows_data)
            item_h  = Mm(block.get('item_h_mm', block.get('h_mm', 50)))
            content_h = item_h * n + GAP_SM * (n - 1)
        else:
            content_h = Mm(block.get('h_mm', 50))
    else:
        content_h = Mm(block.get('h_mm', _DEFAULTS.get(btype, 30)))
    return title_h + content_h + GAP
